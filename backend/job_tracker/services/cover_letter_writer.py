"""The Cover Letter Writer: a cover letter for one job from the facts in my resume and the posting, and the .docx
it's turned into when the form wants a file."""

import io
from dataclasses import dataclass
from typing import Any

import docx
from docx.shared import Pt

from job_tracker.clients.claude import LLM, LLMError
from job_tracker.models.job import Job
from job_tracker.services.fit_scorer import job_text

MAX_OUTPUT_TOKENS = 2_000
MAX_CHARS = 10_000

SYSTEM_PROMPT = """You write a cover letter for one job seeker applying to one job. They will review and edit it \
before it's sent.

Rules:
- Never invent experience, employers, job titles, dates, metrics, education, skills or personal facts. Use only \
what the resume states, and only what the posting says about the company and the role.
- Connect the job seeker's most relevant experience to the posting's main requirements, in the first person, \
plainly and specifically. No clichés, no flattery.
- Three or four short paragraphs, under 350 words, with a greeting and a sign-off with their name as the resume \
gives it. Separate paragraphs with a blank line. Plain text only: no markdown, no placeholders like [Company]."""

SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {"cover_letter": {"type": "string"}},
    "required": ["cover_letter"],
    "additionalProperties": False,
}


@dataclass(frozen=True)
class CoverLetterResult:
    text: str
    model: str
    input_tokens: int
    output_tokens: int


def build_prompt(resume_text: str, job: Job) -> str:
    return (
        f"<resume>\n{resume_text}\n</resume>\n\n<job>\n{job_text(job)}\n</job>\n\nWrite my cover letter for this job."
    )


def tidy(text: str) -> str:
    """The letter with its lines trimmed and at most one blank line between paragraphs."""
    paragraphs = [" ".join(line.split()) for line in text.strip().split("\n")]
    tidied: list[str] = []
    for line in paragraphs:
        if line or (tidied and tidied[-1]):
            tidied.append(line)
    return "\n".join(tidied).strip()


def as_docx(text: str) -> bytes:
    """The letter as a plain Word document: one paragraph per block, keeping the lines within a block."""
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.size = Pt(11)
    for block in text.split("\n\n"):
        paragraph = document.add_paragraph()
        for n, line in enumerate(block.split("\n")):
            if n > 0:
                paragraph.add_run().add_break()
            paragraph.add_run(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class CoverLetterWriter:
    def __init__(self, llm: LLM) -> None:
        self._llm = llm

    def write(self, api_key: str, model: str, resume_text: str, job: Job) -> CoverLetterResult:
        """Raises LLMError if the call fails or the answer doesn't make sense."""
        response = self._llm.generate(
            api_key=api_key,
            model=model,
            system=SYSTEM_PROMPT,
            prompt=build_prompt(resume_text, job),
            schema=SCHEMA,
            max_tokens=MAX_OUTPUT_TOKENS,
        )
        text = response.data.get("cover_letter")
        if not isinstance(text, str) or not tidy(text):
            raise LLMError("Claude's cover letter didn't have the expected fields.")
        return CoverLetterResult(tidy(text)[:MAX_CHARS], response.model, response.input_tokens, response.output_tokens)
