import io
from typing import Any

import docx
import httpx2
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt
from fastapi.testclient import TestClient

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SUMMARY = "Backend engineer with six years of experience building payment systems in Python and Go."
SKILL_LINES = ["Languages: Python, Go, TypeScript, SQL", "Infrastructure: PostgreSQL; Docker; Kubernetes; AWS (Lambda, S3)"]


def sample_resume(summary: str = SUMMARY, pages: int = 1) -> bytes:
    """A resume laid out the way real ones are: a title, headings, labelled skill lines, a table row for a
    job's title and dates, and bullet points. Extra pages repeat the experience section after a page break."""
    document = docx.Document()
    document.add_heading("Jordan Rivera", level=0)
    contact = document.add_paragraph("Brooklyn, NY · jordan@example.com · github.com/jrivera")
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("Summary", level=1)
    document.add_paragraph(summary)
    document.add_heading("Skills", level=1)
    for line in SKILL_LINES:
        document.add_paragraph(line)
    document.add_paragraph()
    for page in range(pages):
        if page > 0:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_heading("Experience" if page == 0 else "Earlier experience", level=1)
        row = document.add_table(rows=1, cols=2).rows[0]
        title = row.cells[0].paragraphs[0].add_run(f"Senior Software Engineer, Initech {page + 1}")
        title.bold = True
        title.font.size = Pt(11)
        row.cells[1].paragraphs[0].text = "2021 – Present"
        for n in range(3):
            document.add_paragraph(f"Built ledger service number {page * 3 + n + 1} handling card payments.", "List Bullet")
    document.add_heading("Education", level=1)
    document.add_paragraph("B.S. Computer Science, State University, 2018")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload(client: TestClient, data: bytes, filename: str = "Jordan Rivera.docx") -> httpx2.Response:
    response: httpx2.Response = client.post("/api/resume", files={"file": (filename, data, DOCX_TYPE)})
    return response


def index_of(resume: dict[str, Any], text: str) -> int:
    return next(int(p["index"]) for p in resume["paragraphs"] if p["text"] == text)


def mapping_for(resume: dict[str, Any], summary: str = SUMMARY) -> dict[str, list[int]]:
    """The sample resume's Summary paragraph and its two skill lines."""
    return {"summary": [index_of(resume, summary)], "skills": [index_of(resume, line) for line in SKILL_LINES]}


def add_ready_resume(client: TestClient, summary: str = SUMMARY) -> dict[str, Any]:
    """Upload the sample resume and mark its Summary and Skills, so it can be scored against."""
    uploaded = upload(client, sample_resume(summary))
    assert uploaded.status_code == 201, uploaded.text
    mapped = client.put("/api/resume/mapping", json=mapping_for(uploaded.json(), summary))
    assert mapped.status_code == 200, mapped.text
    resume: dict[str, Any] = mapped.json()
    return resume
