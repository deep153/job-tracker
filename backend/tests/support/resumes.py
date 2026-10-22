import io
from typing import Any

import docx
import httpx2
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from fastapi.testclient import TestClient

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SUMMARY = "Backend engineer with six years of experience building payment systems in Python and Go."
SKILL_LINES = [
    "Languages: Python, Go, TypeScript, SQL",
    "Infrastructure: PostgreSQL; Docker; Kubernetes; AWS (Lambda, S3)",
]


def sample_resume(summary: str = SUMMARY, pages: int = 1, bold_labels: bool = False) -> bytes:
    """A resume laid out the way real ones are: a title, headings, labelled skill lines, a table row for a
    job's title and dates, and bullet points. Extra pages repeat the experience section after a page break.
    With `bold_labels`, each skill line's label ("Languages:") is a bold run of its own."""
    document = docx.Document()
    document.add_heading("Jordan Rivera", level=0)
    contact = document.add_paragraph("Brooklyn, NY · jordan@example.com · github.com/jrivera")
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("Summary", level=1)
    document.add_paragraph(summary)
    document.add_heading("Skills", level=1)
    for line in SKILL_LINES:
        if bold_labels:
            label, items = line.split(": ", 1)
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{label}: ").bold = True
            paragraph.add_run(items)
        else:
            document.add_paragraph(line)
    document.add_paragraph()
    for page in range(pages):
        if page > 0:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_heading("Experience" if page == 0 else "Earlier experience", level=1)
        row = document.add_table(rows=1, cols=2).rows[0]
        title = row.cells[0].paragraphs[0].add_run(f"Senior Software Engineer, Initech {page + 1}")
        title.bold = True
        title.font.size = Pt(11)
        row.cells[1].paragraphs[0].text = "2021 – Present"
        for n in range(3):
            document.add_paragraph(
                f"Built ledger service number {page * 3 + n + 1} handling card payments.", "List Bullet"
            )
    document.add_heading("Education", level=1)
    document.add_paragraph("B.S. Computer Science, State University, 2018")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


LINKEDIN_URL = "https://www.linkedin.com/in/jordan-rivera-42"
EMAIL = "jordan.rivera@example.com"
PHONE = "(555) 123-4567"
BULLETS = [
    "Payments: Cut checkout latency by 60% for 2M+ users by moving the ledger service to Go.",
    "Reliability: Led on-call for 9+ services and reduced incidents by 35% with Grafana alerts.",
    "Migration: Moved 120 batch jobs from cron to Kubernetes with zero downtime.",
    "Built an internal API used by 14 teams to reconcile card settlements daily.",
]
PUBLICATION = "Rivera, J. Scaling double-entry ledgers. Payments Engineering Journal, 2019."


def _hyperlink(paragraph: Paragraph, text: str, url: str) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    link.append(run)
    paragraph._p.append(link)


def full_resume() -> bytes:
    """A resume shaped like a real one: an all-caps name, a contact line whose LinkedIn URL exists only as a
    hyperlink target, Skills lines with bold labels, list-item experience bullets with bold labels and numbers
    (one without a label), two degrees, and publications that aren't list items."""
    document = docx.Document()
    document.add_heading("JORDAN RIVERA", level=0)
    contact = document.add_paragraph(f"Brooklyn, NY | {PHONE} | {EMAIL} | ")
    _hyperlink(contact, "LinkedIn", LINKEDIN_URL)
    contact.add_run(" | github.com/jrivera")
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("Summary", level=1)
    document.add_paragraph(SUMMARY)
    document.add_heading("Skills", level=1)
    for line in SKILL_LINES:
        label, items = line.split(": ", 1)
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(items)
    document.add_paragraph()
    document.add_heading("Experience", level=1)
    document.add_paragraph().add_run("Senior Software Engineer, Initech | 2021 – Present").bold = True
    for bullet in BULLETS:
        paragraph = document.add_paragraph(style="List Bullet")
        if ": " in bullet:
            label, rest = bullet.split(": ", 1)
            paragraph.add_run(f"{label}: ").bold = True
            paragraph.add_run(rest)
        else:
            paragraph.add_run(bullet)
    document.add_heading("Education", level=1)
    document.add_paragraph("Master of Science in Computer Science, Columbia University | May 2020")
    document.add_paragraph("B.Tech. in Computer Engineering, University of Mumbai | June 2017")
    document.add_heading("Publications", level=1)
    document.add_paragraph(PUBLICATION)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload(client: TestClient, data: bytes, filename: str = "Jordan Rivera.docx") -> httpx2.Response:
    response: httpx2.Response = client.post("/api/resume", files={"file": (filename, data, DOCX_TYPE)})
    return response


def index_of(resume: dict[str, Any], text: str) -> int:
    return next(int(p["index"]) for p in resume["paragraphs"] if p["text"] == text)


def mapping_for(resume: dict[str, Any], summary: str = SUMMARY) -> dict[str, list[int]]:
    """The sample resume's Summary paragraph and its two skill lines, with no experience bullets."""
    return {
        "summary": [index_of(resume, summary)],
        "skills": [index_of(resume, line) for line in SKILL_LINES],
        "bullets": [],
    }


def add_ready_resume(
    client: TestClient,
    summary: str = SUMMARY,
    bold_labels: bool = False,
    data: bytes | None = None,
    bullets: bool = False,
) -> dict[str, Any]:
    """Upload the sample resume (or `data`) and mark its Summary and Skills, so it can be scored against.
    With `bullets`, the proposed experience bullets are marked too."""
    uploaded = upload(client, data or sample_resume(summary, bold_labels=bold_labels))
    assert uploaded.status_code == 201, uploaded.text
    mapping = mapping_for(uploaded.json(), summary)
    mapped = client.put("/api/resume/mapping", json=mapping)
    assert mapped.status_code == 200, mapped.text
    if bullets:
        mapped = client.put("/api/resume/mapping", json={**mapping, "bullets": mapped.json()["proposed_bullets"]})
        assert mapped.status_code == 200, mapped.text
    resume: dict[str, Any] = mapped.json()
    return resume
