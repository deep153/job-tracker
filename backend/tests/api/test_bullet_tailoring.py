import io
from typing import Any

import docx
from fastapi.testclient import TestClient

from job_tracker.services.resume_reader import read_paragraphs
from tests.support.applications import apply, page_count, pdf_text, scored_job
from tests.support.fake_job_boards import FakeJobBoards
from tests.support.fake_llm import FakeLLM
from tests.support.resumes import BULLETS, full_resume

TAILORED = "Backend engineer who builds payments APIs in Python and Go."
LONG = " ".join(["Backend engineer who ships reliable, well-tested payment systems for large card networks."] * 40)
PAYMENTS = "Cut checkout latency by 60% for 2M+ users by moving the ledger service to Go."
REWORDED = "Cut payment API latency by 60% for 2M+ users by moving the ledger service to Go."
API = "Built an internal Python API used by 14 teams to reconcile card settlements daily."


def tailoring(*bullets: tuple[int, str], summary: str = TAILORED) -> dict[str, Any]:
    return {
        "summary": [summary],
        "skills": [],
        "keywords": [],
        "gaps": [],
        "bullets": [{"bullet": number, "text": text} for number, text in bullets],
    }


def bullets_job(client: TestClient, boards: FakeJobBoards) -> int:
    """A job scored against the full resume with its proposed bullets marked; the scorer's keywords for it are
    Python, PostgreSQL and Kafka."""
    return scored_job(client, boards, resume=full_resume(), bullets=True)


def by_label(application: dict[str, Any]) -> dict[str | None, dict[str, Any]]:
    return {bullet["label"]: bullet for bullet in application["resume"]["bullets"]}


def test_reworded_bullets_keep_their_bold_label_and_nothing_else_changes(
    client: TestClient, boards: FakeJobBoards, llm: FakeLLM
) -> None:
    job_id = bullets_job(client, boards)
    llm.tailor(tailoring((1, REWORDED), (4, API)))

    application = apply(client, job_id)

    assert application["status"] == "in_review", application
    bullets = application["resume"]["bullets"]
    assert [(b["label"], b["master"], b["tailored"], b["reverted"]) for b in bullets] == [
        ("Payments", PAYMENTS, REWORDED, None),
        ("Reliability", BULLETS[1].split(": ", 1)[1], BULLETS[1].split(": ", 1)[1], None),
        ("Migration", BULLETS[2].split(": ", 1)[1], BULLETS[2].split(": ", 1)[1], None),
        (None, BULLETS[3], API, None),
    ]

    master_docx = client.get(f"/api/resume/versions/{application['resume']['master_version']}/original.docx").content
    tailored_docx = client.get(f"/api/applications/{application['id']}/resume.docx").content
    mapping = client.get("/api/resume").json()["mapping"]
    changed = {bullets[0]["index"], bullets[3]["index"]}
    master, tailored = read_paragraphs("m.docx", master_docx), read_paragraphs("t.docx", tailored_docx)
    for before, after in zip(master, tailored, strict=True):
        if before.index not in {*mapping["summary"], *mapping["skills"], *changed}:
            assert after == before
    assert tailored[bullets[0]["index"]].text == f"Payments: {REWORDED}"
    assert tailored[bullets[3]["index"]].text == API
    payments = next(p for p in docx.Document(io.BytesIO(tailored_docx)).paragraphs if p.text.startswith("Payments"))
    assert [(run.text, run.bold) for run in payments.runs if run.text] == [("Payments: ", True), (REWORDED, None)]

    pdf = client.get(f"/api/applications/{application['id']}/resume.pdf").content
    assert f"Payments: {REWORDED}" in pdf_text(pdf)
    assert page_count(pdf) == application["resume"]["master_page_count"] == application["resume"]["page_count"]


def test_a_reworded_bullet_that_changes_a_number_is_reverted(
    client: TestClient, boards: FakeJobBoards, llm: FakeLLM
) -> None:
    job_id = bullets_job(client, boards)
    changed = REWORDED.replace("60%", "70%")
    llm.tailor(tailoring((1, changed)))

    application = apply(client, job_id)

    payments = by_label(application)["Payments"]
    assert payments["tailored"] == PAYMENTS
    assert payments["reverted"] == {"text": changed, "reason": "It changes 60% to 70%."}
    text = pdf_text(client.get(f"/api/applications/{application['id']}/resume.pdf").content)
    assert f"Payments: {PAYMENTS}" in text
    assert "70%" not in text


def test_a_reworded_bullet_that_adds_an_unlisted_skill_is_reverted(
    client: TestClient, boards: FakeJobBoards, llm: FakeLLM
) -> None:
    job_id = bullets_job(client, boards)
    with_kafka = "Cut checkout latency by 60% for 2M+ users by moving the ledger to Go and Kafka."
    keeps_grafana = "Led on-call for 9+ services and cut incidents by 35% with Grafana and Python alerts."
    llm.tailor(tailoring((1, with_kafka), (2, keeps_grafana)))

    application = apply(client, job_id)

    payments, reliability = by_label(application)["Payments"], by_label(application)["Reliability"]
    assert payments["tailored"] == PAYMENTS
    assert payments["reverted"] == {
        "text": with_kafka,
        "reason": "It adds “Kafka”, which isn't in this bullet or your master skills list.",
    }
    assert (reliability["tailored"], reliability["reverted"]) == (keeps_grafana, None)


def test_a_reworded_bullet_that_grew_too_much_is_reverted(
    client: TestClient, boards: FakeJobBoards, llm: FakeLLM
) -> None:
    job_id = bullets_job(client, boards)
    longer = f"{PAYMENTS[:-1]}, which made checkout feel instant on mobile."
    llm.tailor(tailoring((1, longer)))

    payments = by_label(apply(client, job_id))["Payments"]

    assert payments["tailored"] == PAYMENTS
    assert payments["reverted"]["text"] == longer
    assert payments["reverted"]["reason"] == "It's 57% longer than the original; a rewrite can grow by 15% at most."


MASTER_SKILLS = [["Python", "Go", "TypeScript", "SQL"], ["PostgreSQL", "Docker", "Kubernetes", "AWS (Lambda, S3)"]]


def edit(client: TestClient, application: dict[str, Any], bullets: list[tuple[int, str]]) -> Any:
    """Regenerate with the Summary and Skills as they are and these bullet texts, by paragraph index."""
    resume = application["resume"]
    return client.put(
        f"/api/applications/{application['id']}/resume",
        json={
            "summary": resume["summary"]["tailored"],
            "skills": [line["skills"] for line in resume["skills"]["tailored"]],
            "bullets": [{"index": index, "text": text} for index, text in bullets],
        },
    )


def test_a_hand_edit_that_changes_a_number_is_reverted(client: TestClient, boards: FakeJobBoards, llm: FakeLLM) -> None:
    application = apply(client, bullets_job(client, boards))
    payments = by_label(application)["Payments"]
    typo = PAYMENTS.replace("2M+", "3M+")

    response = edit(client, application, [(payments["index"], typo)])

    assert response.status_code == 200, response.text
    edited = by_label(response.json())["Payments"]
    assert edited["tailored"] == PAYMENTS
    assert edited["reverted"] == {"text": typo, "reason": "It changes 2M+ to 3M+."}
    assert "3M+" not in pdf_text(client.get(response.json()["resume"]["pdf_url"]).content)


def test_a_valid_hand_edit_is_regenerated_and_rescored(client: TestClient, boards: FakeJobBoards, llm: FakeLLM) -> None:
    application = apply(client, bullets_job(client, boards))
    payments = by_label(application)["Payments"]
    mine = "Cut payments latency by 60% for 2M+ users by moving ledger services to Go."
    llm.score_resume(mine, 93)

    response = edit(client, application, [(payments["index"], mine)])

    assert response.status_code == 200, response.text
    resume = response.json()["resume"]
    assert (by_label(response.json())["Payments"]["tailored"], resume["fit_score"]["after"]) == (mine, 93)
    assert f"Payments: {mine}" in pdf_text(client.get(resume["pdf_url"]).content)
    tailored_docx = client.get(resume["docx_url"]).content
    paragraph = next(p for p in docx.Document(io.BytesIO(tailored_docx)).paragraphs if p.text.startswith("Payments"))
    assert [(run.text, run.bold) for run in paragraph.runs if run.text] == [("Payments: ", True), (mine, None)]
    assert costs_by_kind(client)["rescoring"] == 2
    assert response.json()["cost_usd"] > application["cost_usd"]


def test_use_original_restores_one_bullet_and_keeps_the_others(
    client: TestClient, boards: FakeJobBoards, llm: FakeLLM
) -> None:
    job_id = bullets_job(client, boards)
    llm.tailor(tailoring((1, REWORDED), (4, API)))
    application = apply(client, job_id)
    payments = by_label(application)["Payments"]

    response = edit(client, application, [(payments["index"], payments["master"])])

    bullets = response.json()["resume"]["bullets"]
    assert [(b["tailored"], b["reverted"]) for b in bullets][::3] == [(PAYMENTS, None), (API, None)]
    text = pdf_text(client.get(response.json()["resume"]["pdf_url"]).content)
    assert f"Payments: {PAYMENTS}" in text
    assert API in text


def test_bullet_edits_must_name_a_marked_bullet_and_not_be_blank(
    client: TestClient, boards: FakeJobBoards, llm: FakeLLM
) -> None:
    application = apply(client, bullets_job(client, boards))
    index = by_label(application)["Payments"]["index"]

    for bullets, message in [
        ([(index, "  ")], "Write each bullet, or use the original."),
        ([(0, "Anything")], "Paragraph 0 isn't one of your experience bullets."),
    ]:
        response = edit(client, application, bullets)
        assert response.status_code == 422
        assert response.json()["detail"] == message


def costs_by_kind(client: TestClient) -> dict[str, int]:
    return {k["kind"]: k["calls"] for k in client.get("/api/settings/costs").json()["by_kind"]}


def test_shortening_and_retrying_covers_the_bullets_too(
    client: TestClient, boards: FakeJobBoards, llm: FakeLLM
) -> None:
    job_id = bullets_job(client, boards)
    llm.tailor(tailoring((1, REWORDED), summary=LONG), tailoring((4, API)))

    application = apply(client, job_id)

    assert application["status"] == "in_review", application
    assert len(llm.tailoring_calls()) == 2
    assert [b["tailored"] for b in application["resume"]["bullets"]][::3] == [PAYMENTS, API]
    pdf = client.get(f"/api/applications/{application['id']}/resume.pdf").content
    assert page_count(pdf) == 1
    assert API in pdf_text(pdf)
